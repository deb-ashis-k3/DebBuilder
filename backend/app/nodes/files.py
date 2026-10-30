"""File and system nodes (n8n parity): Extract From File, Convert To File, Read/Write Files from
Disk, Compression, Edit Image, Execute Command, SSH, FTP/SFTP, GraphQL.

Files travel on items as binary references (app/binary.py); ctx.file() / ctx.put_file() move bytes.
"""

from __future__ import annotations

import asyncio
import csv
import fnmatch
import gzip
import io
import json
import os
import subprocess
import zipfile
from pathlib import Path

import httpx

from ..config import get_settings
from .base import NodeContext, NodeError, NodeType, prop, register
from .core import CORE, apply_auth

FILES = dict(group="data", color="#0ea5e9")
SYSTEM = dict(group="core", color="#64748b")


def _with_file(item: dict, prop_name: str, ref: dict, keep_json: bool = True) -> dict:
    base = {k: v for k, v in item.items() if k != "binary"} if keep_json else {}
    return {**base, "binary": {**(item.get("binary") or {}), prop_name: ref}}


# ---------------------------------------------------------------------------
# Extract From File
# ---------------------------------------------------------------------------


def _rows_csv(data: bytes, delimiter: str, header: bool) -> list[dict]:
    text = data.decode("utf-8-sig", "replace")
    if header:
        return [dict(r) for r in csv.DictReader(io.StringIO(text), delimiter=delimiter or ",")]
    return [{f"column_{k}": v for k, v in enumerate(r)} for r in csv.reader(io.StringIO(text), delimiter=delimiter or ",")]


def _rows_xlsx(data: bytes, sheet: str | None, header: bool) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    cell = lambda v: v.isoformat() if hasattr(v, "isoformat") else v  # noqa: E731
    if header:
        keys = [str(h) if h is not None else f"column_{k}" for k, h in enumerate(rows[0])]
        return [{keys[k]: cell(v) for k, v in enumerate(r) if k < len(keys)} for r in rows[1:] if any(v is not None for v in r)]
    return [{f"column_{k}": cell(v) for k, v in enumerate(r)} for r in rows]


def _pdf(data: bytes) -> dict:
    from pypdf import PdfReader

    r = PdfReader(io.BytesIO(data))
    pages = [p.extract_text() or "" for p in r.pages]
    info = {k.lstrip("/"): str(v) for k, v in (r.metadata or {}).items()}
    return {"text": "\n\n".join(pages), "numpages": len(pages), "info": info, "pages": pages}


def _docx(data: bytes) -> str:
    from docx import Document

    return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)


def file_text(data: bytes, ref: dict) -> str:
    """Best-effort plain text of a file (PDF, Word, HTML, anything textual) — used by RAG loaders too."""
    mime, ext = ref.get("mimeType") or "", (ref.get("fileExtension") or "").lower()
    if mime == "application/pdf" or ext == "pdf":
        return _pdf(data)["text"]
    if ext == "docx" or "wordprocessingml" in mime:
        return _docx(data)
    if ext in ("html", "htm") or mime == "text/html":
        from bs4 import BeautifulSoup

        return BeautifulSoup(data, "html.parser").get_text("\n", strip=True)
    if ext in ("xlsx",) or "spreadsheetml" in mime:
        return "\n".join(json.dumps(r, default=str) for r in _rows_xlsx(data, None, True))
    return data.decode("utf-8", "replace")


async def _extract(ctx: NodeContext):
    out = []
    for i, it in enumerate(ctx.items or [{}]):
        p = ctx.params(i)
        op = p.get("operation") or "csv"
        data, ref = ctx.file(it, p.get("binary_property") or "data")
        field = p.get("output_field") or "data"
        try:
            if op == "csv":
                out += _rows_csv(data, p.get("delimiter") or ",", p.get("header", True))
            elif op == "xlsx":
                out += _rows_xlsx(data, p.get("sheet") or None, p.get("header", True))
            elif op == "pdf":
                res = _pdf(data)
                if not p.get("keep_pages"):
                    res.pop("pages")
                out.append({field: res["text"], **{k: v for k, v in res.items() if k != "text"}})
            elif op == "json":
                parsed = json.loads(data.decode("utf-8-sig"))
                out += parsed if isinstance(parsed, list) and p.get("split_arrays", True) else [{field: parsed}]
            elif op == "html":
                from bs4 import BeautifulSoup

                out.append({field: BeautifulSoup(data, "html.parser").get_text("\n", strip=True)})
            elif op == "docx":
                out.append({field: _docx(data)})
            elif op == "base64":
                import base64

                out.append({field: base64.b64encode(data).decode(), "mimeType": ref.get("mimeType"), "fileName": ref.get("fileName")})
            else:  # text
                out.append({field: data.decode(p.get("encoding") or "utf-8", "replace")})
        except NodeError:
            raise
        except Exception as e:  # noqa: BLE001
            raise NodeError(f"Couldn't read '{ref.get('fileName')}' as {op}: {e}", ctx.name) from e
    return {"main": out}


register(NodeType(
    type="file.extract", label="Extract From File", icon="file-output",
    description="Read a file into items: CSV, Excel, PDF, JSON, HTML, Word, plain text, or Base64.",
    properties=[
        prop("operation", "File type", "options", "csv", options=[
            {"label": "CSV", "value": "csv"}, {"label": "Excel (XLSX)", "value": "xlsx"}, {"label": "PDF", "value": "pdf"},
            {"label": "JSON", "value": "json"}, {"label": "HTML (text)", "value": "html"}, {"label": "Word (DOCX)", "value": "docx"},
            {"label": "Text", "value": "text"}, {"label": "Move file to Base64 string", "value": "base64"}]),
        prop("binary_property", "Binary property", "string", "data"),
        prop("header", "First row is the header", "boolean", True, show_if={"operation": ["csv", "xlsx"]}),
        prop("delimiter", "Delimiter", "string", ",", show_if={"operation": ["csv"]}),
        prop("sheet", "Sheet (empty = first)", "string", "", show_if={"operation": ["xlsx"]}),
        prop("keep_pages", "Also output text per page", "boolean", False, show_if={"operation": ["pdf"]}),
        prop("split_arrays", "A JSON array becomes one item per element", "boolean", True, show_if={"operation": ["json"]}),
        prop("encoding", "Encoding", "string", "utf-8", show_if={"operation": ["text"]}),
        prop("output_field", "Output field", "string", "data", show_if={"operation": ["pdf", "html", "docx", "text", "base64", "json"]}),
    ],
    execute=_extract, **FILES,
))

# ---------------------------------------------------------------------------
# Convert To File
# ---------------------------------------------------------------------------


def _flat_rows(items: list[dict]) -> tuple[list[str], list[dict]]:
    rows = [{k: (json.dumps(v, default=str) if isinstance(v, (dict, list)) else v) for k, v in it.items() if k != "binary"} for it in items]
    cols: list[str] = []
    for r in rows:
        cols += [k for k in r if k not in cols]
    return cols, rows


async def _convert(ctx: NodeContext):
    p = ctx.params(0)
    op = p.get("operation") or "csv"
    prop_name = p.get("binary_property") or "data"
    items = [it for it in ctx.items if it]
    name = p.get("file_name") or f"file.{ {'xlsx': 'xlsx', 'json': 'json', 'html': 'html', 'text': 'txt', 'base64': 'bin'}.get(op, 'csv')}"
    if op == "text" or op == "base64":  # one file per item from a field
        out = []
        for i, it in enumerate(ctx.items or [{}]):
            q = ctx.params(i)
            raw = q.get("source_field_value")
            if op == "base64":
                import base64

                data = base64.b64decode(str(raw or ""))
            else:
                data = (raw if isinstance(raw, str) else json.dumps(raw, default=str)).encode()
            out.append(_with_file(it, prop_name, ctx.put_file(data, q.get("file_name") or name, q.get("mime_type") or None)))
        return {"main": out}
    if op == "json":
        if p.get("mode") == "each":
            out = []
            for i, it in enumerate(ctx.items):
                data = json.dumps({k: v for k, v in it.items() if k != "binary"}, indent=2, default=str).encode()
                out.append(_with_file(it, prop_name, ctx.put_file(data, ctx.params(i).get("file_name") or name, "application/json")))
            return {"main": out}
        data = json.dumps([{k: v for k, v in it.items() if k != "binary"} for it in items], indent=2, default=str).encode()
        return {"main": [{"binary": {prop_name: ctx.put_file(data, name, "application/json")}}]}
    cols, rows = _flat_rows(items)
    if op == "csv":
        buf = io.StringIO()
        w = csv.DictWriter(buf, fieldnames=cols, delimiter=p.get("delimiter") or ",", extrasaction="ignore")
        if p.get("header", True):
            w.writeheader()
        w.writerows(rows)
        data, mime = buf.getvalue().encode("utf-8"), "text/csv"
    elif op == "xlsx":
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = (p.get("sheet") or "Sheet1")[:31]
        if p.get("header", True):
            ws.append(cols)
        for r in rows:
            ws.append([r.get(c) for c in cols])
        buf = io.BytesIO()
        wb.save(buf)
        data, mime = buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif op == "html":
        from html import escape

        head = "".join(f"<th>{escape(str(c))}</th>" for c in cols)
        body = "".join("<tr>" + "".join(f"<td>{escape(str(r.get(c, '')))}</td>" for c in cols) + "</tr>" for r in rows)
        data, mime = f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>".encode(), "text/html"
    else:
        raise NodeError(f"Unknown format {op}")
    return {"main": [{"binary": {prop_name: ctx.put_file(data, name, mime)}}]}


register(NodeType(
    type="file.convert", label="Convert To File", icon="file-input",
    description="Turn items into a file: CSV, Excel, JSON, an HTML table, text, or Base64 → file.",
    properties=[
        prop("operation", "Convert to", "options", "csv", options=[
            {"label": "CSV", "value": "csv"}, {"label": "Excel (XLSX)", "value": "xlsx"}, {"label": "JSON", "value": "json"},
            {"label": "HTML table", "value": "html"}, {"label": "Text file (from a field)", "value": "text"},
            {"label": "Base64 string → file", "value": "base64"}]),
        prop("mode", "Mode", "options", "all", options=[{"label": "All items in one file", "value": "all"},
                                                         {"label": "One file per item", "value": "each"}], show_if={"operation": ["json"]}),
        prop("source_field_value", "Content", "text", "{{ $json.data }}", rows=2, show_if={"operation": ["text", "base64"]}),
        prop("mime_type", "MIME type (empty = from file name)", "string", "", show_if={"operation": ["text", "base64"]}),
        prop("header", "Header row", "boolean", True, show_if={"operation": ["csv", "xlsx"]}),
        prop("delimiter", "Delimiter", "string", ",", show_if={"operation": ["csv"]}),
        prop("sheet", "Sheet name", "string", "Sheet1", show_if={"operation": ["xlsx"]}),
        prop("file_name", "File name", "string", ""),
        prop("binary_property", "Put the file in binary property", "string", "data"),
    ],
    execute=_convert, **FILES,
))

# ---------------------------------------------------------------------------
# Read / Write Files from Disk (sandboxed to FILES_ROOT)
# ---------------------------------------------------------------------------


def files_root() -> Path:
    root = Path(get_settings().files_root).resolve()
    root.mkdir(parents=True, exist_ok=True)
    return root


def safe_path(rel: str) -> Path:
    """Resolve a user path inside FILES_ROOT; anything escaping it is refused."""
    root = files_root()
    p = (root / str(rel or "").lstrip("/\\")).resolve()
    if p != root and root not in p.parents:
        raise NodeError(f"Path '{rel}' is outside the allowed folder ({root}). Set FILES_ROOT to change it.")
    return p


async def _read_write(ctx: NodeContext):
    out = []
    root = files_root()
    for i, it in enumerate(ctx.items or [{}]):
        p = ctx.params(i)
        if p.get("operation") == "write":
            data, ref = ctx.file(it, p.get("binary_property") or "data")
            path = safe_path(p.get("path") or ref.get("fileName") or "file")
            path.parent.mkdir(parents=True, exist_ok=True)
            if p.get("append") and path.exists():
                with path.open("ab") as f:
                    f.write(data)
            else:
                path.write_bytes(data)
            out.append({**it, "fileName": str(path.relative_to(root)), "fileSize": len(data)})
            continue
        pattern = str(p.get("path") or "*")
        base = safe_path(os.path.dirname(pattern) or ".")
        matches = sorted(f for f in base.rglob("*") if f.is_file() and fnmatch.fnmatch(str(f.relative_to(root)).replace("\\", "/"), pattern.lstrip("/")))
        if not matches and not p.get("allow_empty"):
            raise NodeError(f"No file matches '{pattern}' in {root}", ctx.name)
        for f in matches[: int(p.get("limit") or 1000)]:
            ref = ctx.put_file(f.read_bytes(), f.name)
            out.append({"fileName": str(f.relative_to(root)).replace("\\", "/"), "binary": {p.get("binary_property") or "data": ref}})
    return {"main": out}


register(NodeType(
    type="file.read_write", label="Read/Write Files from Disk", icon="hard-drive",
    description="Read files matching a pattern, or write a file, inside the server's files folder (FILES_ROOT).",
    properties=[
        prop("operation", "Operation", "options", "read", options=[{"label": "Read file(s)", "value": "read"},
                                                                    {"label": "Write file", "value": "write"}]),
        prop("path", "Path (relative to FILES_ROOT; read accepts * and ** patterns)", "string", "*.csv"),
        prop("binary_property", "Binary property", "string", "data"),
        prop("append", "Append if the file exists", "boolean", False, show_if={"operation": ["write"]}),
        prop("allow_empty", "Don't fail when nothing matches", "boolean", False, show_if={"operation": ["read"]}),
        prop("limit", "Max files", "number", 1000, show_if={"operation": ["read"]}),
    ],
    execute=_read_write, **FILES,
))

# ---------------------------------------------------------------------------
# Compression
# ---------------------------------------------------------------------------


async def _compression(ctx: NodeContext):
    out = []
    for i, it in enumerate(ctx.items or [{}]):
        p = ctx.params(i)
        props = [x.strip() for x in str(p.get("binary_property") or "data").split(",") if x.strip()]
        if p.get("operation") == "decompress":
            data, ref = ctx.file(it, props[0])
            files: dict = {}
            name = ref.get("fileName") or "file"
            if zipfile.is_zipfile(io.BytesIO(data)):
                with zipfile.ZipFile(io.BytesIO(data)) as z:
                    for k, info in enumerate(x for x in z.infolist() if not x.is_dir()):
                        if info.file_size > get_settings().max_binary_mb * 1024 * 1024:
                            raise NodeError(f"'{info.filename}' is larger than MAX_BINARY_MB", ctx.name)
                        files[f"{p.get('output_prefix') or 'file_'}{k}"] = ctx.put_file(z.read(info), Path(info.filename).name)
            elif data[:2] == b"\x1f\x8b":
                files[f"{p.get('output_prefix') or 'file_'}0"] = ctx.put_file(gzip.decompress(data), name.removesuffix(".gz"))
            else:
                raise NodeError(f"'{name}' is not a zip or gzip file", ctx.name)
            out.append({**{k: v for k, v in it.items() if k != "binary"}, "binary": files})
            continue
        fmt = p.get("format") or "zip"
        if fmt == "gzip":
            files = {}
            for pr in props:
                data, ref = ctx.file(it, pr)
                files[pr] = ctx.put_file(gzip.compress(data), f"{ref.get('fileName') or 'file'}.gz", "application/gzip")
            out.append({**{k: v for k, v in it.items() if k != "binary"}, "binary": files})
            continue
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for pr in props:
                data, ref = ctx.file(it, pr)
                z.writestr(ref.get("fileName") or pr, data)
        ref = ctx.put_file(buf.getvalue(), p.get("file_name") or "archive.zip", "application/zip")
        out.append({**{k: v for k, v in it.items() if k != "binary"}, "binary": {p.get("output_property") or "data": ref}})
    return {"main": out}


register(NodeType(
    type="file.compression", label="Compression", icon="file-archive",
    description="Zip or gzip files, or unpack zip / gzip archives.",
    properties=[
        prop("operation", "Operation", "options", "compress", options=["compress", "decompress"]),
        prop("binary_property", "Binary properties (comma separated)", "string", "data"),
        prop("format", "Format", "options", "zip", options=["zip", "gzip"], show_if={"operation": ["compress"]}),
        prop("file_name", "Zip file name", "string", "archive.zip", show_if={"format": ["zip"]}),
        prop("output_property", "Output binary property", "string", "data", show_if={"format": ["zip"]}),
        prop("output_prefix", "Output property prefix", "string", "file_", show_if={"operation": ["decompress"]}),
    ],
    execute=_compression, **FILES,
))

# ---------------------------------------------------------------------------
# Edit Image (Pillow)
# ---------------------------------------------------------------------------


async def _image(ctx: NodeContext):
    from PIL import Image, ImageDraw, ImageFilter, ImageOps

    out = []
    for i, it in enumerate(ctx.items or [{}]):
        p = ctx.params(i)
        prop_name = p.get("binary_property") or "data"
        data, ref = ctx.file(it, prop_name)
        try:
            img = Image.open(io.BytesIO(data))
            img.load()
        except Exception as e:  # noqa: BLE001
            raise NodeError(f"'{ref.get('fileName')}' is not an image Pillow can read: {e}", ctx.name) from e
        op = p.get("operation") or "resize"
        if op == "information":
            out.append({**it, "image": {"format": img.format, "width": img.width, "height": img.height, "mode": img.mode}})
            continue
        if op == "resize":
            w, h = int(p.get("width") or img.width), int(p.get("height") or img.height)
            img = ImageOps.contain(img, (w, h)) if p.get("keep_ratio", True) else img.resize((w, h))
        elif op == "crop":
            x, y = int(p.get("x") or 0), int(p.get("y") or 0)
            img = img.crop((x, y, x + int(p.get("width") or img.width), y + int(p.get("height") or img.height)))
        elif op == "rotate":
            img = img.rotate(-float(p.get("degrees") or 90), expand=True)
        elif op == "blur":
            img = img.filter(ImageFilter.GaussianBlur(float(p.get("radius") or 2)))
        elif op == "grayscale":
            img = ImageOps.grayscale(img)
        elif op == "text":
            img = img.convert("RGBA")
            ImageDraw.Draw(img).text((int(p.get("x") or 10), int(p.get("y") or 10)), str(p.get("text") or ""),
                                     fill=str(p.get("color") or "#ffffff"))
        fmt = (p.get("format") or (img.format or ref.get("fileExtension") or "png")).upper().replace("JPG", "JPEG")
        if fmt == "JPEG" and img.mode in ("RGBA", "P", "LA"):
            img = img.convert("RGB")
        buf = io.BytesIO()
        img.save(buf, format=fmt)
        name = f"{Path(ref.get('fileName') or 'image').stem}.{fmt.lower().replace('jpeg', 'jpg')}"
        out.append(_with_file(it, prop_name, ctx.put_file(buf.getvalue(), name, f"image/{fmt.lower()}")))
    return {"main": out}


register(NodeType(
    type="file.edit_image", label="Edit Image", icon="image",
    description="Resize, crop, rotate, blur, grayscale, add text, convert format, or read image information.",
    properties=[
        prop("operation", "Operation", "options", "resize", options=["resize", "crop", "rotate", "blur", "grayscale", "text", "information"]),
        prop("binary_property", "Binary property", "string", "data"),
        prop("width", "Width", "number", 800, show_if={"operation": ["resize", "crop"]}),
        prop("height", "Height", "number", 600, show_if={"operation": ["resize", "crop"]}),
        prop("keep_ratio", "Keep aspect ratio", "boolean", True, show_if={"operation": ["resize"]}),
        prop("x", "X", "number", 0, show_if={"operation": ["crop", "text"]}),
        prop("y", "Y", "number", 0, show_if={"operation": ["crop", "text"]}),
        prop("degrees", "Degrees (clockwise)", "number", 90, show_if={"operation": ["rotate"]}),
        prop("radius", "Blur radius", "number", 2, show_if={"operation": ["blur"]}),
        prop("text", "Text", "string", "", show_if={"operation": ["text"]}),
        prop("color", "Text colour", "string", "#ffffff", show_if={"operation": ["text"]}),
        prop("format", "Output format (empty = same)", "options", "", options=["", "png", "jpeg", "webp", "gif", "bmp"]),
    ],
    execute=_image, **FILES,
))

# ---------------------------------------------------------------------------
# Execute Command (off unless ENABLE_EXECUTE_COMMAND=true)
# ---------------------------------------------------------------------------


async def _command(ctx: NodeContext):
    if not get_settings().enable_execute_command:
        raise NodeError("Execute Command is disabled on this server. An admin can enable it with ENABLE_EXECUTE_COMMAND=true "
                        "(it runs shell commands with the server's permissions).", ctx.name)
    out = []
    targets = range(len(ctx.items or [{}])) if ctx.param("each") else [0]
    for i in targets:
        cmd = str(ctx.param("command", i) or "")
        timeout = float(ctx.param("timeout", i) or 60)
        # A thread, like the Code node: the Windows selector loop (needed by psycopg) has no subprocess support.
        try:
            proc = await asyncio.to_thread(subprocess.run, cmd, shell=True, capture_output=True, timeout=timeout, cwd=str(files_root()))
        except subprocess.TimeoutExpired as e:
            raise NodeError(f"Command timed out after {timeout:g}s", ctx.name) from e
        res = {"exitCode": proc.returncode, "stdout": proc.stdout.decode("utf-8", "replace").rstrip(),
               "stderr": proc.stderr.decode("utf-8", "replace").rstrip()}
        if proc.returncode and not ctx.param("ignore_errors", i):
            raise NodeError(f"Command failed (exit {proc.returncode}): {res['stderr'][:500]}", ctx.name)
        out.append(res)
    return {"main": out}


register(NodeType(
    type="system.command", label="Execute Command", icon="terminal",
    description="Run a shell command on the server (disabled unless the admin sets ENABLE_EXECUTE_COMMAND=true).",
    properties=[prop("command", "Command", "code", "echo hello", rows=3), prop("each", "Run once for each item", "boolean", False),
                prop("timeout", "Timeout (seconds)", "number", 60), prop("ignore_errors", "Don't fail on a non-zero exit code", "boolean", False)],
    execute=_command, **SYSTEM,
))

# ---------------------------------------------------------------------------
# SSH (commands + SFTP) and FTP
# ---------------------------------------------------------------------------


def _ssh_opts(c: dict) -> dict:
    import asyncssh

    opts: dict = {"host": c["host"], "port": int(c.get("port") or 22), "username": c.get("username") or None,
                  "known_hosts": None, "connect_timeout": 20}
    if c.get("private_key"):
        opts["client_keys"] = [asyncssh.import_private_key(c["private_key"], c.get("passphrase") or None)]
    if c.get("password"):
        opts["password"] = c["password"]
    return opts


async def _ssh(ctx: NodeContext):
    import asyncssh

    _, c = await ctx.credential()
    out = []
    async with asyncssh.connect(**_ssh_opts(c)) as conn:
        for i, it in enumerate(ctx.items or [{}]):
            p = ctx.params(i)
            op = p.get("operation") or "execute"
            if op == "execute":
                r = await conn.run(str(p.get("command") or ""), check=False, timeout=float(p.get("timeout") or 60))
                out.append({"stdout": str(r.stdout or "").rstrip(), "stderr": str(r.stderr or "").rstrip(), "exitCode": r.exit_status})
                continue
            async with conn.start_sftp_client() as sftp:
                path = str(p.get("path") or "")
                if op == "download":
                    async with sftp.open(path, "rb") as f:
                        data = await f.read()
                    out.append(_with_file(it, p.get("binary_property") or "data", ctx.put_file(data, os.path.basename(path))))
                elif op == "upload":
                    data, ref = ctx.file(it, p.get("binary_property") or "data")
                    target = path if not path.endswith("/") else path + (ref.get("fileName") or "file")
                    async with sftp.open(target, "wb") as f:
                        await f.write(data)
                    out.append({**it, "path": target, "fileSize": len(data)})
                elif op == "list":
                    for name in await sftp.listdir(path or "."):
                        if name not in (".", ".."):
                            out.append({"name": name, "path": f"{path.rstrip('/')}/{name}" if path else name})
                elif op == "delete":
                    await sftp.remove(path)
                    out.append({**it, "deleted": path})
                elif op == "rename":
                    await sftp.rename(path, str(p.get("new_path") or ""))
                    out.append({**it, "from": path, "to": p.get("new_path")})
    return {"main": out}


SSH_OPS = [{"label": "Execute command", "value": "execute"}, {"label": "Download file (SFTP)", "value": "download"},
           {"label": "Upload file (SFTP)", "value": "upload"}, {"label": "List folder (SFTP)", "value": "list"},
           {"label": "Delete file (SFTP)", "value": "delete"}, {"label": "Rename / move (SFTP)", "value": "rename"}]
register(NodeType(
    type="net.ssh", label="SSH", icon="square-terminal",
    description="Run commands on a remote server, or move files over SFTP.",
    properties=[
        prop("credential", "Credential", "credential", None, credential_types=["ssh"], required=True),
        prop("operation", "Operation", "options", "execute", options=SSH_OPS),
        prop("command", "Command", "code", "uptime", rows=3, show_if={"operation": ["execute"]}),
        prop("timeout", "Timeout (seconds)", "number", 60, show_if={"operation": ["execute"]}),
        prop("path", "Remote path", "string", "", show_if={"operation": ["download", "upload", "list", "delete", "rename"]}),
        prop("new_path", "New path", "string", "", show_if={"operation": ["rename"]}),
        prop("binary_property", "Binary property", "string", "data", show_if={"operation": ["download", "upload"]}),
    ],
    execute=_ssh, **SYSTEM,
))


async def _ftp(ctx: NodeContext):
    import ftplib

    ctype, c = await ctx.credential()
    if ctype == "ssh":  # SFTP through the SSH node's implementation
        return await _ssh(ctx)
    items = ctx.items or [{}]
    params = [ctx.params(i) for i in range(len(items))]
    files = {i: ctx.file(items[i], params[i].get("binary_property") or "data") for i, p in enumerate(params) if p.get("operation") == "upload"}

    def run() -> list:
        f = ftplib.FTP_TLS(timeout=30) if c.get("tls") else ftplib.FTP(timeout=30)
        f.connect(c["host"], int(c.get("port") or 21))
        f.login(c.get("username") or "anonymous", c.get("password") or "")
        if c.get("tls"):
            f.prot_p()
        res = []
        try:
            for i, p in enumerate(params):
                op, path = p.get("operation") or "list", str(p.get("path") or "")
                if op == "list":
                    for name, facts in f.mlsd(path or "."):
                        if name not in (".", ".."):
                            res.append(("row", {"name": name, "type": facts.get("type"), "size": int(facts.get("size") or 0)}))
                elif op == "download":
                    buf = io.BytesIO()
                    f.retrbinary(f"RETR {path}", buf.write)
                    res.append(("file", (i, buf.getvalue(), os.path.basename(path))))
                elif op == "upload":
                    data, ref = files[i]
                    target = path if not path.endswith("/") else path + (ref.get("fileName") or "file")
                    f.storbinary(f"STOR {target}", io.BytesIO(data))
                    res.append(("row", {**items[i], "path": target, "fileSize": len(data)}))
                elif op == "delete":
                    f.delete(path)
                    res.append(("row", {**items[i], "deleted": path}))
                elif op == "rename":
                    f.rename(path, str(p.get("new_path") or ""))
                    res.append(("row", {**items[i], "from": path, "to": p.get("new_path")}))
        finally:
            try:
                f.quit()
            except Exception:  # noqa: BLE001
                f.close()
        return res

    out = []
    for kind, v in await asyncio.to_thread(run):
        if kind == "file":
            i, data, name = v
            out.append(_with_file(items[i], params[i].get("binary_property") or "data", ctx.put_file(data, name)))
        else:
            out.append(v)
    return {"main": out}


register(NodeType(
    type="net.ftp", label="FTP / SFTP", icon="folder-sync",
    description="List, download, upload, delete or rename files on an FTP/FTPS server (or SFTP with an SSH credential).",
    properties=[
        prop("credential", "Credential", "credential", None, credential_types=["ftp", "ssh"], required=True),
        prop("operation", "Operation", "options", "list", options=["list", "download", "upload", "delete", "rename"]),
        prop("path", "Remote path", "string", "/"),
        prop("new_path", "New path", "string", "", show_if={"operation": ["rename"]}),
        prop("binary_property", "Binary property", "string", "data", show_if={"operation": ["download", "upload"]}),
    ],
    execute=_ftp, **SYSTEM,
))

# ---------------------------------------------------------------------------
# GraphQL
# ---------------------------------------------------------------------------


async def _graphql(ctx: NodeContext):
    out = []
    for i in range(len(ctx.items or [{}])):
        p = ctx.params(i)
        headers = {"Content-Type": "application/json", **{h["name"]: str(h.get("value") or "") for h in p.get("headers") or [] if h.get("name")}}
        auth, params = await apply_auth(ctx, i, headers)
        variables = p.get("variables")
        variables = json.loads(variables) if isinstance(variables, str) and variables.strip() else (variables or {})
        body = {"query": p.get("query") or "", "variables": variables, **({"operationName": p["operation_name"]} if p.get("operation_name") else {})}
        async with httpx.AsyncClient(timeout=float(p.get("timeout") or 30)) as c:
            r = await c.post(p["endpoint"], json=body, headers=headers, params=params or None, auth=auth)
        try:
            data = r.json()
        except ValueError as e:
            raise NodeError(f"GraphQL endpoint returned non-JSON ({r.status_code}): {r.text[:300]}", ctx.name) from e
        if (r.status_code >= 400 or data.get("errors")) and not p.get("continue_on_errors"):
            raise NodeError(f"GraphQL error: {json.dumps(data.get('errors') or data)[:500]}", ctx.name)
        out.append(data)
    return {"main": out}


register(NodeType(
    type="core.graphql", label="GraphQL", icon="share-2",
    description="Send a GraphQL query or mutation (any credential type for auth).",
    properties=[
        prop("endpoint", "Endpoint", "string", "https://", required=True),
        prop("query", "Query", "code", "query {\n  viewer { login }\n}", rows=8),
        prop("variables", "Variables (JSON)", "json", "{}"),
        prop("operation_name", "Operation name", "string", ""),
        prop("auth", "Authentication", "options", "none", options=[{"label": "None", "value": "none"},
                                                                    {"label": "Credential", "value": "credential"}]),
        prop("credential", "Credential", "credential", None, show_if={"auth": ["credential"]}),
        prop("headers", "Headers", "collection", [], fields=[prop("name", "Name"), prop("value", "Value", "string", "")]),
        prop("continue_on_errors", "Return GraphQL errors instead of failing", "boolean", False),
        prop("timeout", "Timeout (seconds)", "number", 30),
    ],
    execute=_graphql, **CORE,
))
